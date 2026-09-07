"""Generate deterministic Phase 9 report QA artifacts without touching plant data or email."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.api import service


PLANT_TIMEZONE = ZoneInfo("Asia/Calcutta")
METER_ONE = {"meter_id": "MTR-001", "meter_name": "Screen Printing", "location": "OSP"}
METER_TWO = {"meter_id": "MTR-002", "meter_name": "TC Oven", "location": "OSP"}


def _export(filters: dict, format_name: str, rows_by_meter: dict[str, list[dict]]) -> dict:
    def fetch_rows(_connection, meter_id, _parameter_keys, _start, _end):
        return rows_by_meter[meter_id]

    with (
        patch("app.api.service._require_known_meters", return_value=[METER_ONE, METER_TWO]),
        patch("app.api.service._open_connection") as open_connection,
        patch("app.api.service._fetch_report_rows", side_effect=fetch_rows),
    ):
        open_connection.return_value.__enter__.return_value = object()
        return service.build_export_payload(filters, format_name)


def _scheduled_export(rows: list[dict]) -> dict:
    with (
        patch("app.api.service._require_known_meters", return_value=[METER_ONE]),
        patch("app.api.service._open_connection") as open_connection,
        patch("app.api.service._fetch_report_source_rows", return_value=rows),
    ):
        open_connection.return_value.__enter__.return_value = object()
        return service.build_scheduled_report_payload(
            meter_ids=["MTR-001"],
            parameter_keys=["active_power_total", "active_energy_received_out_of_load"],
            reading_time_text="08:00",
            start=datetime(2026, 8, 21, 0, 0, tzinfo=PLANT_TIMEZONE),
            end=datetime(2026, 8, 21, 8, 5, tzinfo=PLANT_TIMEZONE),
            interval_hours=None,
            window_mode="start_to_current",
        )


def _assert_excel(path: Path, expected_charts: int) -> dict:
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, f"{path.name} has an invalid XLSX ZIP member"
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    assert "Graph Data" in workbook.sheetnames, f"{path.name} is missing Graph Data"
    assert len(sheet._charts) == expected_charts, f"{path.name} has {len(sheet._charts)} charts, expected {expected_charts}"
    return {
        "filename": path.name,
        "sheets": workbook.sheetnames,
        "charts": len(sheet._charts),
        "data_rows": max(0, sheet.max_row - 2),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="Directory for generated QA evidence")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    start = datetime(2026, 8, 21, 8, 0, tzinfo=PLANT_TIMEZONE)
    end = start + timedelta(hours=2)
    meter_one_rows = [
        {"timestamp": start, "active_power_total": 0.0, "active_energy_received_out_of_load": 100.0},
        {"timestamp": start + timedelta(hours=1), "active_power_total": 2.5, "active_energy_received_out_of_load": 102.0},
        {"timestamp": end, "active_power_total": 3.0, "active_energy_received_out_of_load": 105.0},
        {"timestamp": end + timedelta(minutes=1), "active_power_total": None, "active_energy_received_out_of_load": None},
    ]
    meter_two_rows = [
        {"timestamp": start, "active_power_total": 4.0, "active_energy_received_out_of_load": 200.0},
        {"timestamp": start + timedelta(hours=1), "active_power_total": 5.0, "active_energy_received_out_of_load": 204.0},
        {"timestamp": end, "active_power_total": 6.0, "active_energy_received_out_of_load": 207.0},
    ]
    filters = {
        "meterIds": ["MTR-001", "MTR-002"],
        "parameterKeys": ["active_power_total", "active_energy_received_out_of_load"],
        "startDateTime": start.isoformat(),
        "endDateTime": end.isoformat(),
        "intervalHours": None,
    }
    rows_by_meter = {"MTR-001": meter_one_rows, "MTR-002": meter_two_rows}
    all_readings = _export(filters, "xlsx", rows_by_meter)
    all_readings_path = args.output / "multi_all_readings.xlsx"
    all_readings_path.write_bytes(all_readings["bytes"])

    hourly = _export({**filters, "intervalHours": 1}, "xlsx", rows_by_meter)
    hourly_path = args.output / "multi_hourly.xlsx"
    hourly_path.write_bytes(hourly["bytes"])

    word = _export(filters, "docx", rows_by_meter)
    word_path = args.output / "multi_all_readings.docx"
    word_path.write_bytes(word["bytes"])

    scheduled_rows = [
        {"timestamp": start - timedelta(minutes=10), "active_power_total": 1.0, "active_energy_received_out_of_load": 300.0},
        {"timestamp": start - timedelta(minutes=2), "active_power_total": 2.0, "active_energy_received_out_of_load": 301.0},
        {"timestamp": start + timedelta(minutes=3), "active_power_total": 3.0, "active_energy_received_out_of_load": 302.0},
    ]
    scheduled = _scheduled_export(scheduled_rows)
    scheduled_path = args.output / "scheduled_snapshot.xlsx"
    scheduled_path.write_bytes(scheduled["bytes"])

    evidence = {
        "multi_all_readings": _assert_excel(all_readings_path, expected_charts=2),
        "multi_hourly": _assert_excel(hourly_path, expected_charts=2),
        "scheduled_snapshot": _assert_excel(scheduled_path, expected_charts=2),
        "zero_value_preserved": load_workbook(all_readings_path, data_only=True)["Graph Data"].cell(row=2, column=2).value == 0.0,
        "timestamp_only_rows_removed": all_readings["rows"] == 6,
        "scheduled_target_time": load_workbook(scheduled_path, data_only=True).active.cell(row=3, column=2).value.strftime("%H:%M"),
        "scheduled_actual_snapshot_time": load_workbook(scheduled_path, data_only=True)["Graph Data"].cell(row=2, column=1).value.strftime("%H:%M"),
    }
    with zipfile.ZipFile(word_path) as archive:
        document_xml = archive.read("word/document.xml")
    evidence["word_table_grid_present"] = b"<w:tblGrid>" in document_xml
    assert evidence["zero_value_preserved"], "valid zero value disappeared from graph data"
    assert evidence["timestamp_only_rows_removed"], "timestamp-only report row was not removed"
    assert evidence["scheduled_target_time"] == "08:00", "scheduled target time is incorrect"
    assert evidence["scheduled_actual_snapshot_time"] == "07:58", "scheduled snapshot did not prefer the latest earlier reading"
    assert evidence["word_table_grid_present"], "Word table grid is missing"

    (args.output / "summary.json").write_text(json.dumps(evidence, indent=2), encoding="utf-8")
    print(json.dumps(evidence, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
